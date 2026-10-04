import json
from io import BytesIO
from datetime import date
from decimal import Decimal

import app.models
import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.api.v1.admin.backup import router as admin_backup_router
from app.core.enums import AccountRole, TourType, VehicleType
from app.models.account import Account
from app.models.base import Base
from app.models.customer_profile import CustomerProfile
from app.models.destination import Destination
from app.models.hotel import Hotel
from app.models.tour_detail import TourDetail
from app.models.tour_departure import TourDeparture
from app.models.tour_package import TourPackage
from app.models.tour_variant import TourVariant
from app.models.vehicle import Vehicle
from app.models.vendor import Vendor
from app.services.backup_service import BackupService
from app.schemas.response import ActionResponse


@pytest.fixture
def db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(bind=engine)
    with Session(engine) as session:
        yield session
    Base.metadata.drop_all(bind=engine)
    engine.dispose()


def add_backup_records(db: Session) -> None:
    destination = Destination(name="Darjeeling", slug="darjeeling-backup", country="India")
    customer = Account(
        account_code="C-BACKUP-1",
        name="Backup Customer",
        role=AccountRole.CUSTOMER,
        email="backup-customer@example.com",
    )
    staff = Account(
        account_code="S-BACKUP-1",
        name="Backup Staff",
        role=AccountRole.STAFF,
        email="backup-staff@example.com",
    )
    admin = Account(
        account_code="A-BACKUP-1",
        name="Backup Admin",
        role=AccountRole.ADMIN,
        email="backup-admin@example.com",
    )
    db.add_all([destination, customer, staff, admin])
    db.flush()
    customer_profile = CustomerProfile(
        account_id=customer.id,
        referral_code="BACKUP-CUSTOMER",
        address="Coochbehar",
    )
    package = TourPackage(
        tour_code="T-BACKUP-1",
        slug="darjeeling-backup-tour",
        title="Darjeeling Tour",
        destination_id=destination.id,
        type=TourType.DOMESTIC,
    )
    db.add_all([customer_profile, package])
    db.flush()
    variant = TourVariant(
        package_id=package.id,
        slug="darjeeling-backup-tour-standard",
        name="Standard",
        valid_from=date(2026, 1, 1),
        valid_to=date(2026, 12, 31),
        duration_days=3,
        duration_nights=2,
        list_price=Decimal("15000.00"),
        selling_price=Decimal("12500.00"),
    )
    db.add(variant)
    db.flush()
    db.add(
        TourDeparture(
            variant_id=variant.id,
            departure_date=date(2026, 10, 17),
            return_date=date(2026, 10, 25),
            total_seats=20,
            available_seats=12,
        )
    )
    db.add_all(
        [
            TourDetail(
                variant_id=variant.id,
                banner={"url": "banner.jpg"},
                gallery=[],
                highlights=[{"title": "Tea gardens"}],
                inclusions=["Breakfast"],
                exclusions=[],
                itinerary=[{"day": 1, "title": "Arrival"}],
                route_stops=[],
            ),
            Hotel(name="Backup Hotel", image=[], destination_id=destination.id),
            Vendor(name="Backup Vendor", type="HOTEL", email="vendor@example.com"),
            Vehicle(
                name="Backup Vehicle",
                vehicle_image=[],
                vehicle_type=VehicleType.ANY,
                capacity=4,
                price_per_day=Decimal("2000.00"),
            ),
        ]
    )
    db.commit()


def test_import_endpoint_uses_action_response():
    import_route = next(
        route for route in admin_backup_router.routes if route.path.endswith("/import")
    )
    assert import_route.response_model is ActionResponse


@pytest.mark.parametrize("format", ["json", "csv", "xlsx"])
def test_backup_round_trip_preserves_selected_relationships(db: Session, format: str):
    add_backup_records(db)
    groups = ["customers", "staff", "tours", "hotels", "vendors", "vehicles"]
    content, _, _ = BackupService(db).export(groups, format)
    if format == "json":
        manifest = json.loads(content)
        assert manifest["tables"]["tour_departures"] == [
            {
                "variant_id": str(db.scalar(select(TourVariant)).id),
                "departure_date": "2026-10-17",
                "return_date": "2026-10-25",
                "total_seats": 20,
                "available_seats": 12,
                "id": manifest["tables"]["tour_departures"][0]["id"],
                "created_at": manifest["tables"]["tour_departures"][0]["created_at"],
                "updated_at": manifest["tables"]["tour_departures"][0]["updated_at"],
                "is_active": True,
            }
        ]
    if format == "xlsx":
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        active_sheet = workbook.active
        rows = active_sheet.iter_rows(values_only=True)
        assert active_sheet.title == "tour_overview"
        columns = next(rows)
        record = dict(zip(columns, next(rows)))
        assert record["tour_title"] == "Darjeeling Tour"
        assert record["destination_name"] == "Darjeeling"
        assert record["variant_name"] == "Standard"
        assert json.loads(record["itinerary"]) == [{"day": 1, "title": "Arrival"}]
        workbook.close()

    target_engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(bind=target_engine)
    with Session(target_engine) as target:
        imported = BackupService(target).import_backup(content, format)
        assert imported == {"inserted": 11, "updated": 0}
        assert target.scalar(select(Account).where(Account.role == AccountRole.ADMIN)) is None
        assert target.scalar(select(CustomerProfile)).account.name == "Backup Customer"
        detail = target.scalar(select(TourDetail))
        assert detail.variant.package.destination.name == "Darjeeling"
        assert detail.itinerary == [{"day": 1, "title": "Arrival"}]
        departure = target.scalar(select(TourDeparture))
        assert departure.departure_date == date(2026, 10, 17)
        assert departure.return_date == date(2026, 10, 25)
        assert departure.total_seats == 20
        assert departure.available_seats == 12
        assert target.scalar(select(Hotel)).destination_ref.name == "Darjeeling"

        second_import = BackupService(target).import_backup(content, format)
        assert second_import == {"inserted": 0, "updated": 11}
    target_engine.dispose()


def test_backup_import_rejects_admin_account_records_without_partial_write(db: Session):
    admin = Account(
        account_code="A-IMPORT-1",
        name="Must Not Import",
        role=AccountRole.ADMIN,
        email="must-not-import@example.com",
    )
    db.add(admin)
    db.flush()
    manifest = {
        "format": "ct-admin-backup",
        "version": 1,
        "groups": ["staff"],
        "tables": {"accounts": [BackupService._serialize_model(admin)]},
    }

    with pytest.raises(ValueError, match="Admin accounts cannot be imported"):
        BackupService(db).import_backup(json.dumps(manifest).encode(), "json")
    assert db.scalar(select(Account)) is None