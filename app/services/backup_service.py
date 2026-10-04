import csv
import io
import json
import uuid
import zipfile
from datetime import date, datetime, timezone
from decimal import Decimal
from enum import Enum
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy.types import JSON

from app.core.enums import AccountRole
from app.core.messages.error import BackupError
from app.core.validation import (
    deserialize_backup_row,
    validate_backup_account_role,
    validate_backup_manifest,
)
from app.models.account import Account
from app.models.customer_profile import CustomerProfile
from app.models.destination import Destination
from app.models.hotel import Hotel
from app.models.tour_detail import TourDetail
from app.models.tour_departure import TourDeparture
from app.models.tour_package import TourPackage
from app.models.tour_variant import TourVariant
from app.models.vehicle import Vehicle
from app.models.vendor import Vendor
from app.repository.backup_repo import BackupRepository


BACKUP_FORMAT = "ct-admin-backup"
BACKUP_VERSION = 1
MAX_BACKUP_SIZE = 25 * 1024 * 1024
MAX_RECORDS = 100_000
TABLE_MODELS = {
    model.__tablename__: model
    for model in (
        Account,
        CustomerProfile,
        Destination,
        Hotel,
        TourPackage,
        TourVariant,
        TourDeparture,
        TourDetail,
        Vendor,
        Vehicle,
    )
}
IMPORT_ORDER = (
    "destinations",
    "accounts",
    "customer_profiles",
    "tour_packages",
    "tour_variants",
    "tour_departures",
    "tour_details",
    "hotels",
    "vendors",
    "vehicles",
)
GROUP_TABLES = {
    "customers": {"accounts", "customer_profiles"},
    "staff": {"accounts"},
    "tours": {
        "tour_packages",
        "tour_variants",
        "tour_departures",
        "tour_details",
        "destinations",
    },
    "destinations": {"destinations"},
    "hotels": {"hotels", "destinations"},
    "vendors": {"vendors"},
    "vehicles": {"vehicles"},
}
VALID_GROUPS = set(GROUP_TABLES)


class BackupService:
    def __init__(self, db):
        self.repo = BackupRepository(db)
        self.db = db

    def export(self, groups: list[str], format: str) -> tuple[bytes, str, str]:
        normalized_groups = sorted(set(groups))
        if not normalized_groups or not set(normalized_groups) <= VALID_GROUPS:
            raise ValueError("Select one or more supported backup groups")
        if format not in {"json", "csv", "xlsx"}:
            raise ValueError("Backup format must be json, csv, or xlsx")

        table_names = set().union(*(GROUP_TABLES[group] for group in normalized_groups))
        tables: dict[str, list[dict[str, Any]]] = {}
        for table_name in IMPORT_ORDER:
            if table_name not in table_names:
                continue
            if table_name == "accounts":
                rows = []
                if "customers" in normalized_groups:
                    rows.extend(self.repo.list_rows(Account, role=AccountRole.CUSTOMER))
                if "staff" in normalized_groups:
                    rows.extend(self.repo.list_rows(Account, role=AccountRole.STAFF))
            elif table_name == "customer_profiles" and "customers" in normalized_groups:
                account_ids = [
                    account.id
                    for account in self.repo.list_rows(Account, role=AccountRole.CUSTOMER)
                ]
                rows = self.repo.list_rows(
                    CustomerProfile,
                    filters={"account_id": account_ids},
                )
            else:
                rows = self.repo.list_rows(TABLE_MODELS[table_name])
            tables[table_name] = [self._serialize_model(row) for row in rows]

        manifest = {
            "format": BACKUP_FORMAT,
            "version": BACKUP_VERSION,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "groups": normalized_groups,
            "tables": tables,
        }
        timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        if format == "json":
            return (
                json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8"),
                f"admin-backup-{timestamp}.json",
                "application/json",
            )
        if format == "csv":
            return (
                self._create_csv_archive(manifest),
                f"admin-backup-{timestamp}.zip",
                "application/zip",
            )
        return (
            self._create_xlsx_workbook(manifest),
            f"admin-backup-{timestamp}.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def import_backup(self, content: bytes, format: str) -> dict[str, int]:
        try:
            if format == "json":
                manifest = json.loads(content.decode("utf-8"))
            elif format == "csv":
                manifest = self._read_csv_archive(content)
            elif format == "xlsx":
                manifest = self._read_xlsx_workbook(content)
            else:
                raise ValueError("Backup format must be json, csv, or xlsx")
            tables, groups = validate_backup_manifest(
                manifest,
                format_name=BACKUP_FORMAT,
                version=BACKUP_VERSION,
                valid_groups=VALID_GROUPS,
                group_tables=GROUP_TABLES,
            )
            if sum(len(rows) for rows in tables.values()) > MAX_RECORDS:
                raise ValueError("Backup contains too many records")
            result = {"inserted": 0, "updated": 0}
            seen_ids: set[tuple[str, Any]] = set()
            imported_ids: dict[tuple[str, Any], Any] = {}
            for table_name in IMPORT_ORDER:
                rows = tables.get(table_name, [])
                model = TABLE_MODELS.get(table_name)
                if rows and model is None:
                    raise ValueError(f"Unsupported backup table: {table_name}")
                for raw_row in rows:
                    values = deserialize_backup_row(model, raw_row)
                    row_id = (table_name, values["id"])
                    if row_id in seen_ids:
                        continue
                    seen_ids.add(row_id)
                    for column in model.__table__.columns:
                        for foreign_key in column.foreign_keys:
                            referenced_table = foreign_key.column.table.name
                            reference = values.get(column.name)
                            reference_key = (referenced_table, reference)
                            if reference_key in imported_ids:
                                values[column.name] = imported_ids[reference_key]
                    if model is Account:
                        validate_backup_account_role(
                            values,
                            groups,
                            customer_role=AccountRole.CUSTOMER,
                            staff_role=AccountRole.STAFF,
                            admin_role=AccountRole.ADMIN,
                        )
                    if model is CustomerProfile:
                        account = self.repo.get_by_id(Account, values["account_id"])
                        if account is None or account.role != AccountRole.CUSTOMER:
                            raise ValueError("Customer profile references a missing customer account")
                    inserted, record = self.repo.insert_if_missing(model, values)
                    imported_ids[row_id] = record.id
                    if inserted:
                        result["inserted"] += 1
            self.db.commit()
            return result
        except ValueError:
            self.db.rollback()
            raise
        except (TypeError, KeyError, UnicodeDecodeError) as exc:
            self.db.rollback()
            raise ValueError("Backup contains invalid or malformed data") from exc
        except Exception:
            self.db.rollback()
            raise

    @staticmethod
    def _serialize_model(model: Any) -> dict[str, Any]:
        return {
            column.name: BackupService._json_value(getattr(model, column.name))
            for column in model.__table__.columns
        }

    @staticmethod
    def _json_value(value: Any) -> Any:
        if isinstance(value, Enum):
            return value.value
        if isinstance(value, (uuid.UUID, Decimal, date, datetime)):
            return value.isoformat() if isinstance(value, (date, datetime)) else str(value)
        if isinstance(value, list):
            return [BackupService._json_value(item) for item in value]
        if isinstance(value, dict):
            return {key: BackupService._json_value(item) for key, item in value.items()}
        return value

    @staticmethod
    def _create_csv_archive(manifest: dict[str, Any]) -> bytes:
        archive_bytes = io.BytesIO()
        tables = manifest["tables"]
        manifest_copy = {**manifest, "tables": sorted(tables)}
        with zipfile.ZipFile(archive_bytes, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("manifest.json", json.dumps(manifest_copy, ensure_ascii=False))
            for table_name, rows in tables.items():
                output = io.StringIO(newline="")
                columns = list(rows[0]) if rows else [
                    column.name for column in TABLE_MODELS[table_name].__table__.columns
                ]
                writer = csv.DictWriter(output, fieldnames=columns)
                writer.writeheader()
                for row in rows:
                    writer.writerow({
                        field: json.dumps(row.get(field), ensure_ascii=False, separators=(",", ":"))
                        for field in columns
                    })
                archive.writestr(f"{table_name}.csv", output.getvalue().encode("utf-8"))
        return archive_bytes.getvalue()

    @staticmethod
    def _create_xlsx_workbook(manifest: dict[str, Any]) -> bytes:
        workbook = Workbook()
        manifest_sheet = workbook.active
        manifest_sheet.title = "manifest"
        manifest_sheet.append(["key", "value"])
        for key in ("format", "version", "created_at", "groups"):
            manifest_sheet.append([key, json.dumps(manifest[key], ensure_ascii=False)])
        table_names = sorted(manifest["tables"])
        manifest_sheet.append(["tables", json.dumps(table_names)])

        tour_overview = None
        if "tours" in manifest["groups"]:
            tour_overview = workbook.create_sheet("tour_overview")
            BackupService._write_tour_overview(tour_overview, manifest["tables"])

        first_populated_sheet = None
        first_table_sheet = None
        for table_name, rows in manifest["tables"].items():
            sheet = workbook.create_sheet(table_name)
            if first_table_sheet is None:
                first_table_sheet = sheet
            if rows and first_populated_sheet is None:
                first_populated_sheet = sheet
            columns = list(rows[0]) if rows else [
                column.name for column in TABLE_MODELS[table_name].__table__.columns
            ]
            sheet.append(columns)
            for row in rows:
                cells = [
                    json.dumps(value, ensure_ascii=False, separators=(",", ":"))
                    if isinstance(value := row.get(column), (dict, list))
                    else value
                    for column in columns
                ]
                if any(isinstance(value, str) and len(value) > 32767 for value in cells):
                    raise ValueError(BackupError.XLSX_FIELD_TOO_LARGE)
                sheet.append(cells)
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            active_sheet = tour_overview or first_populated_sheet or first_table_sheet or manifest_sheet
            workbook.active = workbook.index(active_sheet)

        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _write_tour_overview(sheet: Any, tables: dict[str, list[dict[str, Any]]]) -> None:
        columns = (
            "tour_code",
            "tour_title",
            "tour_slug",
            "destination_name",
            "description",
            "variant_name",
            "variant_slug",
            "season_name",
            "valid_from",
            "valid_to",
            "duration_days",
            "duration_nights",
            "list_price",
            "selling_price",
            "badge",
            "is_default",
            "banner",
            "gallery",
            "highlights",
            "inclusions",
            "exclusions",
            "itinerary",
            "route_stops",
        )
        sheet.append(columns)
        destinations = {
            row["id"]: row.get("name")
            for row in tables.get("destinations", [])
        }
        details = {
            row["variant_id"]: row
            for row in tables.get("tour_details", [])
        }
        variants_by_package: dict[str, list[dict[str, Any]]] = {}
        for variant in tables.get("tour_variants", []):
            variants_by_package.setdefault(variant["package_id"], []).append(variant)

        detail_fields = columns[16:]
        for package in tables.get("tour_packages", []):
            variants = variants_by_package.get(package["id"], []) or [None]
            for variant in variants:
                detail = details.get(variant["id"], {}) if variant else {}
                row = {
                    "tour_code": package.get("tour_code"),
                    "tour_title": package.get("title"),
                    "tour_slug": package.get("slug"),
                    "destination_name": destinations.get(package.get("destination_id")),
                    "description": package.get("description"),
                    "variant_name": variant.get("name") if variant else None,
                    "variant_slug": variant.get("slug") if variant else None,
                    "season_name": variant.get("season_name") if variant else None,
                    "valid_from": variant.get("valid_from") if variant else None,
                    "valid_to": variant.get("valid_to") if variant else None,
                    "duration_days": variant.get("duration_days") if variant else None,
                    "duration_nights": variant.get("duration_nights") if variant else None,
                    "list_price": variant.get("list_price") if variant else None,
                    "selling_price": variant.get("selling_price") if variant else None,
                    "badge": variant.get("badge") if variant else None,
                    "is_default": variant.get("is_default") if variant else None,
                    **{field: detail.get(field) for field in detail_fields},
                }
                cell_values = []
                for column in columns:
                    value = row.get(column)
                    if isinstance(value, (dict, list)):
                        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
                    if isinstance(value, str) and len(value) > 32767:
                        raise ValueError(BackupError.XLSX_FIELD_TOO_LARGE)
                    cell_values.append(value)
                sheet.append(cell_values)
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"

    @staticmethod
    def _read_csv_archive(content: bytes) -> dict[str, Any]:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                expanded_size = archive.getinfo("manifest.json").file_size
                if expanded_size > MAX_BACKUP_SIZE:
                    raise ValueError("CSV backup archive exceeds the expanded size limit")
                manifest = json.loads(archive.read("manifest.json").decode("utf-8"))
                table_names = manifest.get("tables")
                if not isinstance(table_names, list):
                    raise ValueError("CSV backup manifest has no table list")
                if not all(isinstance(table_name, str) for table_name in table_names):
                    raise ValueError("CSV backup manifest contains an invalid table name")
                if len(set(table_names)) != len(table_names):
                    raise ValueError("CSV backup manifest contains duplicate tables")
                tables = {}
                for table_name in table_names:
                    if table_name not in TABLE_MODELS:
                        raise ValueError(f"Unsupported backup table: {table_name}")
                    expanded_size += archive.getinfo(f"{table_name}.csv").file_size
                    if expanded_size > MAX_BACKUP_SIZE:
                        raise ValueError("CSV backup archive contains an oversized table")
                    with archive.open(f"{table_name}.csv") as source:
                        text_source = io.TextIOWrapper(source, encoding="utf-8", newline="")
                        reader = csv.DictReader(text_source)
                        tables[table_name] = [
                            {field: json.loads(value) for field, value in row.items()}
                            for row in reader
                        ]
                manifest["tables"] = tables
                return manifest
        except (zipfile.BadZipFile, KeyError, json.JSONDecodeError, UnicodeDecodeError) as exc:
            raise ValueError("Invalid CSV backup archive") from exc

    @staticmethod
    def _read_xlsx_workbook(content: bytes) -> dict[str, Any]:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > MAX_BACKUP_SIZE:
                    raise ValueError("XLSX backup exceeds the expanded size limit")
            workbook = load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=True,
                keep_links=False,
            )
            try:
                if "manifest" not in workbook.sheetnames:
                    raise ValueError("XLSX backup has no manifest sheet")
                manifest_rows = workbook["manifest"].iter_rows(values_only=True)
                if next(manifest_rows, None) != ("key", "value"):
                    raise ValueError("XLSX backup manifest is invalid")
                manifest = {
                    key: json.loads(value)
                    for key, value in manifest_rows
                    if isinstance(key, str) and isinstance(value, str)
                }
                table_names = manifest.get("tables")
                if not isinstance(table_names, list) or not all(
                    isinstance(table_name, str) for table_name in table_names
                ):
                    raise ValueError("XLSX backup has an invalid table list")
                if len(set(table_names)) != len(table_names):
                    raise ValueError("XLSX backup has duplicate tables")
                tables = {}
                for table_name in table_names:
                    if table_name not in TABLE_MODELS or table_name not in workbook.sheetnames:
                        raise ValueError(f"Unsupported XLSX backup table: {table_name}")
                    rows = workbook[table_name].iter_rows(values_only=True)
                    columns = next(rows, None)
                    if not columns or not all(isinstance(column, str) for column in columns):
                        raise ValueError(f"XLSX backup table {table_name} has invalid columns")
                    if len(set(columns)) != len(columns):
                        raise ValueError(f"XLSX backup table {table_name} has duplicate columns")
                    model_columns = {
                        column.name: column.type
                        for column in TABLE_MODELS[table_name].__table__.columns
                    }
                    table_rows = []
                    for row in rows:
                        record = {}
                        for column, value in zip(columns, row):
                            if isinstance(model_columns.get(column), JSON) and isinstance(value, str):
                                value = json.loads(value)
                            record[column] = value
                        table_rows.append(record)
                    tables[table_name] = table_rows
                manifest["tables"] = tables
                return manifest
            finally:
                workbook.close()
        except (zipfile.BadZipFile, KeyError, json.JSONDecodeError, UnicodeDecodeError) as exc:
            raise ValueError("Invalid XLSX backup workbook") from exc

